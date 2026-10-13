"""Builds Wealth_Training_Data.xlsx: a fictional wealth-management dataset
for the 'From Spreadsheets to Dashboards' ebook. All names and figures are invented."""
import random
from datetime import date, timedelta
import calendar

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

random.seed(42)
np.random.seed(42)

OUT = "Wealth_Training_Data.xlsx"

# ---------- Calendar: 36 month-ends, Oct 2023 to Sep 2026 ----------
def month_end(y, m):
    return date(y, m, calendar.monthrange(y, m)[1])

months = []
y, m = 2023, 10
for _ in range(36):
    months.append(month_end(y, m))
    m += 1
    if m == 13:
        y, m = y + 1, 1
START = months[0]

# ---------- Advisers ----------
first = ["Amelia", "Oliver", "Isla", "George", "Ava", "Harry", "Mia", "Jack", "Sophie", "Noah",
         "Grace", "Leo", "Chloe", "Arthur", "Ella", "Oscar", "Freya", "Henry", "Ruby", "Theo"]
last = ["Bennett", "Clarke", "Dawson", "Ellis", "Fletcher", "Grant", "Hughes", "Irving", "Jarvis", "Kemp",
        "Lowe", "Mason", "Norris", "Osborne", "Pearce", "Quinn", "Rowe", "Sutton", "Turner", "Vaughan"]
random.shuffle(last)
teams = ["Private Clients", "Family Office", "Retirement Planning", "Professionals"]
offices = {"Private Clients": "London", "Family Office": "London",
           "Retirement Planning": "Bristol", "Professionals": "Edinburgh"}
advisers = []
for i in range(20):
    team = teams[i % 4]
    advisers.append({
        "AdviserID": f"A{i+1:02d}",
        "AdviserName": f"{first[i]} {last[i]}",
        "AdviserEmail": f"{first[i].lower()}.{last[i].lower()}@example.com",
        "Team": team,
        "Office": offices[team],
    })
adv_df = pd.DataFrame(advisers)

# ---------- Clients ----------
cfirst = ["Harper", "Owen", "Priya", "Daniel", "Lucy", "Samuel", "Hannah", "Joseph", "Emily", "Thomas",
          "Zara", "William", "Charlotte", "James", "Aisha", "Benjamin", "Olivia", "Lewis", "Megan", "Ethan",
          "Rosie", "Callum", "Imogen", "Adam", "Nina", "Ryan", "Holly", "Marcus", "Erin", "Kieran",
          "Leah", "Felix", "Anya", "Joel", "Tara", "Rhys", "Maya", "Liam", "Esme", "Dylan"]
clast = ["Lane", "Marsh", "Nair", "Wood", "Patel", "Brooks", "Reid", "Price", "Kaur", "Shaw",
         "Ahmed", "Cole", "Ward", "Hart", "Khan", "Fox", "Barker", "Doyle", "Ross", "Bell",
         "Webb", "Lloyd", "Chen", "Hale", "Moss", "Booth", "Frost", "Okafor", "Burke", "Hayes",
         "Pike", "Walsh", "Dunn", "Forde", "Iqbal", "Hunt", "Rees", "Stone", "Carr", "Young"]
regions = ["London", "South East", "South West", "Midlands", "North West", "North East", "Scotland", "Wales"]
region_w = [0.28, 0.20, 0.10, 0.10, 0.10, 0.07, 0.10, 0.05]
segments = ["Mass Affluent", "High Net Worth", "Ultra High Net Worth"]
seg_w = [0.55, 0.37, 0.08]
seg_aum = {"Mass Affluent": (60_000, 250_000), "High Net Worth": (250_000, 2_000_000),
           "Ultra High Net Worth": (2_000_000, 12_000_000)}

clients = []
used = set()
for i in range(500):
    while True:
        name = f"{random.choice(cfirst)} {random.choice(clast)}"
        if name not in used:
            used.add(name)
            break
    seg = random.choices(segments, seg_w)[0]
    # ~80% joined before the data window, ~20% during it
    if random.random() < 0.8:
        jd = date(2012, 1, 1) + timedelta(days=random.randint(0, (START - date(2012, 1, 1)).days - 1))
    else:
        jd = START + timedelta(days=random.randint(1, (months[-1] - START).days - 40))
    clients.append({
        "ClientID": f"C{i+1:04d}",
        "ClientName": name,
        "Region": random.choices(regions, region_w)[0],
        "Segment": seg,
        "AdviserID": random.choice(adv_df["AdviserID"]),
        "JoinDate": jd,
    })

# ---------- Funds and returns ----------
funds = [
    ("Equities", "UK Equity Income Fund", "GBP"),
    ("Equities", "Global Equity Growth Fund", "USD"),
    ("Equities", "European Equity Fund", "EUR"),
    ("Bonds", "UK Gilt Fund", "GBP"),
    ("Bonds", "Global Corporate Bond Fund", "USD"),
    ("Cash", "Sterling Money Market Fund", "GBP"),
    ("Property", "UK Commercial Property Fund", "GBP"),
    ("Alternatives", "Multi-Strategy Absolute Return Fund", "GBP"),
]
mu = {"Equities": 0.0065, "Bonds": 0.0025, "Cash": 0.0035, "Property": 0.0030, "Alternatives": 0.0040}
sd = {"Equities": 0.035, "Bonds": 0.012, "Cash": 0.0005, "Property": 0.018, "Alternatives": 0.010}
ac_ret = {ac: np.random.normal(mu[ac], sd[ac], 36) for ac in mu}
# a common market shock in two months for realism
for k, shock in [(6, -0.045), (19, -0.03)]:
    ac_ret["Equities"][k] += shock
    ac_ret["Property"][k] += shock / 2

# FX: start levels, random walk (GBP per unit of currency)
fx = {"GBP": np.ones(36)}
for cur, start in [("USD", 0.79), ("EUR", 0.86)]:
    fx[cur] = start * np.cumprod(1 + np.random.normal(0, 0.012, 36))

# ---------- Simulate holdings and transactions ----------
holdings, txns = [], []
hid = tid = 0
fee_rate_annual = {"Mass Affluent": 0.0110, "High Net Worth": 0.0085, "Ultra High Net Worth": 0.0065}

def rand_day(month_end_date):
    return date(month_end_date.year, month_end_date.month, random.randint(1, month_end_date.day))

leavers = set(random.sample(range(500), 12))   # clients who withdraw most assets
for idx, c in enumerate(clients):
    seg = c["Segment"]
    nfunds = random.randint(3, 6)
    chosen = random.sample(range(len(funds)), nfunds)
    w = np.random.dirichlet(np.ones(nfunds) * 2)
    lo, hi = seg_aum[seg]
    start_aum = float(np.exp(np.random.uniform(np.log(lo), np.log(hi))))
    active = c["JoinDate"] < START
    vals = None  # value in GBP per chosen fund
    leave_month = random.randint(14, 32) if idx in leavers else None

    for t, me in enumerate(months):
        dep = wd = 0.0
        if not active:
            # joins this month?
            if c["JoinDate"] <= me:
                active = True
                dep = round(start_aum, -2)
                txns.append((c["ClientID"], c["JoinDate"], "Deposit", dep))
                vals = w * dep
            else:
                continue
        else:
            if vals is None:
                vals = w * start_aum
            # market move
            for j, f in enumerate(chosen):
                vals[j] *= 1 + ac_ret[funds[f][0]][t]
            total = vals.sum()
            # deposits
            if random.random() < 0.10:
                dep = round(total * random.uniform(0.02, 0.15), -2)
                txns.append((c["ClientID"], rand_day(me), "Deposit", dep))
            # withdrawals
            if leave_month is not None and t == leave_month:
                wd = round(total * random.uniform(0.6, 0.85), -2)
                txns.append((c["ClientID"], rand_day(me), "Withdrawal", wd))
            elif random.random() < 0.07:
                wd = round(total * random.uniform(0.01, 0.08), -2)
                txns.append((c["ClientID"], rand_day(me), "Withdrawal", wd))
            # quarterly fee (Dec, Mar, Jun, Sep)
            fee = 0.0
            if me.month in (3, 6, 9, 12):
                fee = round(total * fee_rate_annual[seg] / 4, 2)
                txns.append((c["ClientID"], me, "Fee", fee))
            net = dep - wd - fee
            vals = vals + w * net
            vals = np.maximum(vals, 0)

        for j, f in enumerate(chosen):
            ac, fund, cur = funds[f]
            local = vals[j] / fx[cur][t]
            hid += 1
            holdings.append({
                "HoldingID": f"H{hid:06d}",
                "ClientID": c["ClientID"],
                "AssetClass": ac,
                "Fund": fund,
                "Currency": cur,
                "Value": round(float(local), 2),
                "ValuationDate": me,
            })

hold_df = pd.DataFrame(holdings)

txns.sort(key=lambda r: (r[1], r[0]))
tx_rows = []
for cid, d, typ, amt in txns:
    tid += 1
    tx_rows.append({
        "TransactionID": f"T{tid:06d}",
        "ClientID": cid,
        # Stored as UK-format text on purpose: Chapter 2 converts it with "Using locale"
        "Date": d.strftime("%d/%m/%Y"),
        "Type": typ,
        "Amount": round(float(amt), 2),
    })
tx_df = pd.DataFrame(tx_rows)

fx_rows = []
for t, me in enumerate(months):
    for cur in ("GBP", "USD", "EUR"):
        fx_rows.append({"Date": me, "Currency": cur, "RateToGBP": round(float(fx[cur][t]), 4)})
fx_df = pd.DataFrame(fx_rows)

# Benchmark: blended 55% equities, 30% bonds, 10% cash, 5% property, plus noise
bench = (0.55 * ac_ret["Equities"] + 0.30 * ac_ret["Bonds"] + 0.10 * ac_ret["Cash"]
         + 0.05 * ac_ret["Property"] + np.random.normal(0, 0.002, 36))
bm_df = pd.DataFrame({"Date": months, "MonthlyReturn": np.round(bench, 4)})

# ---------- Deliberate untidiness in Clients for Chapter 2 ----------
cl_df = pd.DataFrame(clients)
messy_idx = random.sample(range(500), 40)
for k, i in enumerate(messy_idx):
    r = cl_df.at[i, "Region"]
    variant = [r.lower(), r + " ", r.upper(), " " + r][k % 4]
    cl_df.at[i, "Region"] = variant
dups = cl_df.sample(8, random_state=7)
cl_df = pd.concat([cl_df, dups]).sort_values("ClientID").reset_index(drop=True)

# ---------- Write ----------
with pd.ExcelWriter(OUT, engine="openpyxl", date_format="DD/MM/YYYY") as xw:
    cl_df.to_excel(xw, sheet_name="Clients", index=False)
    adv_df.to_excel(xw, sheet_name="Advisers", index=False)
    hold_df.to_excel(xw, sheet_name="Holdings", index=False)
    tx_df.to_excel(xw, sheet_name="Transactions", index=False)
    fx_df.to_excel(xw, sheet_name="FXRates", index=False)
    bm_df.to_excel(xw, sheet_name="Benchmark", index=False)

wb = load_workbook(OUT)
hdr_fill = PatternFill("solid", start_color="1F3864")
widths = {"ClientName": 20, "AdviserName": 18, "AdviserEmail": 30, "Fund": 36, "Team": 20,
          "Segment": 20, "Region": 14, "AssetClass": 13}
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial", size=10)
    for cell in ws[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = widths.get(cell.value, 15)
    ws.freeze_panes = "A2"
    headers = [c.value for c in ws[1]]
    for name, fmt in [("Value", "#,##0.00"), ("Amount", "#,##0.00"),
                      ("RateToGBP", "0.0000"), ("MonthlyReturn", "0.00%")]:
        if name in headers:
            col = headers.index(name) + 1
            for r in range(2, ws.max_row + 1):
                ws.cell(r, col).number_format = fmt
for name, ws in [("Clients", wb["Clients"])]:
    pass
wb.save(OUT)

print("Clients", len(cl_df), "Advisers", len(adv_df), "Holdings", len(hold_df),
      "Transactions", len(tx_df), "FXRates", len(fx_df), "Benchmark", len(bm_df))
